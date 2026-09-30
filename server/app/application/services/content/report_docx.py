"""
Рендер отчёта по выборке постов (PostReportData) в DOCX — см. post_report.py
(данные) и use_case/posts/summary_report.py (вызывающая ручка,
POST /api/v1/posts/summary/report?fmt=docx).

Cyrillic в DOCX не требует регистрации шрифтов (в отличие от PDF/reportlab) —
рендерится тем приложением, которым открывают файл (Word/LibreOffice/Google
Docs), у них уже есть системные шрифты с кириллицей.
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt

from app.application.services.content.post_report import PostReportData

_SOURCE_TYPE_LABELS = {"vk": "VK", "rss": "RSS", "telegram": "Telegram"}


def render_report_docx(data: PostReportData) -> bytes:
    doc = Document()

    doc.add_heading(data.title or "Отчёт по выборке новостей", level=0)
    meta = doc.add_paragraph()
    meta.add_run(
        f"Сгенерировано: {data.generated_at:%d.%m.%Y %H:%M} UTC · "
        f"Постов в выборке: {data.posts_requested} · "
        f"Учтено в аналитике (с текстом): {data.posts_used} · "
        f"Словарь ЦКМ: {data.lang.value}"
    ).italic = True
    if data.missing_post_ids:
        warn = doc.add_paragraph()
        warn.add_run(f"Не найдены в базе: {', '.join(str(i) for i in data.missing_post_ids)}").italic = True

    if data.summary:
        doc.add_heading("Саммари", level=1)
        doc.add_paragraph(data.summary)
        if data.topics:
            topics_p = doc.add_paragraph()
            topics_p.add_run("Темы: ").bold = True
            topics_p.add_run(", ".join(data.topics))

    if data.schwartz:
        doc.add_heading("ЦКМ-аналитика (агрегат по выборке)", level=1)
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Параметр"
        hdr[1].text = "Доля"
        for param, value in sorted(data.schwartz.items(), key=lambda kv: -kv[1]):
            row = table.add_row().cells
            row[0].text = param
            row[1].text = f"{value * 100:.1f}%"

    if data.categories:
        doc.add_heading("Категории лемм (топ по выборке)", level=1)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Категория"
        hdr[1].text = "Доля"
        for cat, value in list(data.categories.items())[:10]:
            row = table.add_row().cells
            row[0].text = cat
            row[1].text = f"{value * 100:.1f}%"

    if data.posts:
        doc.add_heading("Посты в выборке", level=1)
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, label in enumerate(["ID", "Источник", "Тип", "Дата", "Текст / ссылка"]):
            hdr[i].text = label
        for item in data.posts:
            row = table.add_row().cells
            row[0].text = str(item.id)
            row[1].text = item.source_name or "—"
            row[2].text = _SOURCE_TYPE_LABELS.get(item.source_type or "", item.source_type or "—")
            row[3].text = f"{item.published_at:%d.%m.%Y %H:%M}" if item.published_at else "—"
            text_cell = row[4]
            text_cell.text = item.text_snippet or "(текст отсутствует)"
            if item.url:
                p = text_cell.add_paragraph()
                run = p.add_run(item.url)
                run.font.size = Pt(8)
                run.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
